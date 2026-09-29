#!/usr/bin/env python3
"""
Simplified NSE Stock Scanner with Telegram Integration
Standalone implementation without complex dependencies
"""

import pandas as pd
import yfinance as yf
import os
import sys
from datetime import datetime
import pytz
from concurrent.futures import ThreadPoolExecutor, as_completed
import logging
import json

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# NSE F&O Stock Universe
COMPLETE_NSE_FO_UNIVERSE = [
    '360ONE.NS', 'ABB.NS', 'APLAPOLLO.NS', 'AUBANK.NS', 'ADANIENSOL.NS',
    'ADANIENT.NS', 'ADANIGREEN.NS', 'ADANIPORTS.NS', 'ABCAPITAL.NS', 'ALKEM.NS',
    'AMBER.NS', 'AMBUJACEM.NS', 'ANGELONE.NS', 'APOLLOHOSP.NS', 'ASHOKLEY.NS',
    'ASIANPAINT.NS', 'ASTRAL.NS', 'AUROPHARMA.NS', 'DMART.NS', 'AXISBANK.NS',
    'BSE.NS', 'BAJAJ-AUTO.NS', 'BAJFINANCE.NS', 'BAJAJFINSV.NS', 'BAJAJHLDNG.NS',
    'BANDHANBNK.NS', 'BANKBARODA.NS', 'BANKINDIA.NS', 'BDL.NS', 'BEL.NS',
    'BHARATFORG.NS', 'BHEL.NS', 'BPCL.NS', 'BHARTIARTL.NS', 'BIOCON.NS',
    'BLUESTARCO.NS', 'BOSCHLTD.NS', 'BRITANNIA.NS', 'CGPOWER.NS', 'CANBK.NS',
    'CDSL.NS', 'CHOLAFIN.NS', 'CIPLA.NS', 'COALINDIA.NS', 'COFORGE.NS',
    'COLPAL.NS', 'CAMS.NS', 'CONCOR.NS', 'CROMPTON.NS', 'CUMMINSIND.NS',
    'DLF.NS', 'DABUR.NS', 'DALBHARAT.NS', 'DELHIVERY.NS', 'DIVISLAB.NS',
    'DIXON.NS', 'DRREDDY.NS', 'ETERNAL.NS', 'EICHERMOT.NS', 'EXIDEIND.NS',
    'NYKAA.NS', 'FORTIS.NS', 'GAIL.NS', 'GMRAIRPORT.NS', 'GLENMARK.NS',
    'GODREJCP.NS', 'GODREJPROP.NS', 'GRASIM.NS', 'HCLTECH.NS', 'HDFCAMC.NS',
    'HDFCBANK.NS', 'HDFCLIFE.NS', 'HAVELLS.NS', 'HEROMOTOCO.NS', 'HINDALCO.NS',
    'HAL.NS', 'HINDPETRO.NS', 'HINDUNILVR.NS', 'HINDZINC.NS', 'POWERINDIA.NS',
    'HUDCO.NS', 'ICICIBANK.NS', 'ICICIGI.NS', 'ICICIPRULI.NS', 'IDFCFIRSTB.NS',
    'IIFL.NS', 'ITC.NS', 'INDIANB.NS', 'IEX.NS', 'IOC.NS', 'IRCTC.NS',
    'IRFC.NS', 'IREDA.NS', 'INDUSTOWER.NS', 'INDUSINDBK.NS', 'NAUKRI.NS',
    'INFY.NS', 'INOXWIND.NS', 'INDIGO.NS', 'JINDALSTEL.NS', 'JSWENERGY.NS',
    'JSWSTEEL.NS', 'JIOFIN.NS', 'JUBLFOOD.NS', 'KEI.NS', 'KPITTECH.NS',
    'KALYANKJIL.NS', 'KAYNES.NS', 'KFINTECH.NS', 'KOTAKBANK.NS', 'LTF.NS',
    'LICHSGFIN.NS', 'LTIM.NS', 'LT.NS', 'LAURUSLABS.NS', 'LICI.NS', 'LODHA.NS',
    'LUPIN.NS', 'M&M.NS', 'MANAPPURAM.NS', 'MANKIND.NS', 'MARICO.NS',
    'MARUTI.NS', 'MFSL.NS', 'MAXHEALTH.NS', 'MAZDOCK.NS', 'MPHASIS.NS',
    'MCX.NS', 'MUTHOOTFIN.NS', 'NBCC.NS', 'NHPC.NS', 'NMDC.NS', 'NTPC.NS',
    'NATIONALUM.NS', 'NESTLEIND.NS', 'NUVAMA.NS', 'OBEROIRLTY.NS', 'ONGC.NS',
    'OIL.NS', 'PAYTM.NS', 'OFSS.NS', 'POLICYBZR.NS', 'PGEL.NS', 'PIIND.NS',
    'PNBHOUSING.NS', 'PAGEIND.NS', 'PATANJALI.NS', 'PERSISTENT.NS', 'PETRONET.NS',
    'PIDILITIND.NS', 'PPLPHARMA.NS', 'POLYCAB.NS', 'PFC.NS', 'POWERGRID.NS',
    'PREMIERENE.NS', 'PRESTIGE.NS', 'PNB.NS', 'RBLBANK.NS', 'RECLTD.NS',
    'RVNL.NS', 'RELIANCE.NS', 'SBICARD.NS', 'SBILIFE.NS', 'SHREECEM.NS',
    'SRF.NS', 'SAMMAANCAP.NS', 'MOTHERSON.NS', 'SHRIRAMFIN.NS', 'SIEMENS.NS',
    'SOLARINDS.NS', 'SONACOMS.NS', 'SBIN.NS', 'SAIL.NS', 'SUNPHARMA.NS',
    'SUPREMEIND.NS', 'SUZLON.NS', 'SWIGGY.NS', 'SYNGENE.NS', 'TATACONSUM.NS',
    'TVSMOTOR.NS', 'TCS.NS', 'TATAELXSI.NS', 'TMPV.NS', 'TATAPOWER.NS',
    'TATASTEEL.NS', 'TATATECH.NS', 'TECHM.NS', 'FEDERALBNK.NS', 'INDHOTEL.NS',
    'PHOENIXLTD.NS', 'TITAN.NS', 'TORNTPHARM.NS', 'TORNTPOWER.NS', 'TRENT.NS',
    'TIINDIA.NS', 'UNOMINDA.NS', 'UPL.NS', 'ULTRACEMCO.NS', 'UNIONBANK.NS',
    'UNITDSPR.NS', 'VBL.NS', 'VEDL.NS', 'IDEA.NS', 'VOLTAS.NS', 'WAAREEENER.NS',
    'WIPRO.NS', 'YESBANK.NS', 'ZYDUSLIFE.NS'
]


class SimpleStockScanner:
    """Simple stock scanner for finding trading opportunities"""

    def __init__(self, telegram_token=None, telegram_chat_id=None):
        self.bot_token = telegram_token or os.getenv('TELEGRAM_BOT_TOKEN')
        self.chat_id = telegram_chat_id or os.getenv('TELEGRAM_CHAT_ID')
        self.ist = pytz.timezone('Asia/Kolkata')

        if self.bot_token and self.chat_id:
            self.telegram_enabled = True
            logger.info("✅ Telegram integration enabled")
        else:
            self.telegram_enabled = False
            logger.info("⚠️  Telegram credentials not found - results will be saved to Excel only")

    def send_telegram(self, message):
        """Send message to Telegram"""
        if not self.telegram_enabled:
            return False

        try:
            import requests
            url = f"https://api.telegram.org/bot{self.bot_token}/sendMessage"
            payload = {
                'chat_id': self.chat_id,
                'text': message,
                'parse_mode': 'Markdown'
            }
            response = requests.post(url, json=payload, timeout=10)
            if response.status_code == 200:
                logger.info("📤 Telegram message sent")
                return True
            else:
                logger.error(f"Failed to send Telegram: {response.status_code}")
                return False
        except Exception as e:
            logger.error(f"Telegram error: {e}")
            return False

    def get_stock_data(self, symbol, period="3mo"):
        """Fetch stock data"""
        try:
            data = yf.download(symbol, period=period, progress=False, interval="1d")
            if data is None or len(data) < 20:
                return None
            return data
        except Exception as e:
            logger.debug(f"Error fetching {symbol}: {e}")
            return None

    def calculate_rsi(self, prices, period=14):
        """Calculate RSI"""
        try:
            deltas = prices.diff()
            seed = deltas[:period+1]
            up = seed[seed >= 0].sum() / period
            down = -seed[seed < 0].sum() / period
            rs = up / down if down != 0 else 0
            rsi = 100 - 100 / (1.0 + rs)

            rsis = [None] * (period)
            rsis.append(rsi)

            for i in range(period + 1, len(prices)):
                delta = deltas[i]
                if delta > 0:
                    upval = delta
                    downval = 0.0
                else:
                    upval = 0.0
                    downval = -delta

                up = (up * (period - 1) + upval) / period
                down = (down * (period - 1) + downval) / period
                rs = up / down if down != 0 else 0
                rsi = 100 - 100 / (1.0 + rs)
                rsis.append(rsi)

            return pd.Series(rsis, index=prices.index)
        except:
            return pd.Series([None] * len(prices), index=prices.index)

    def scan_stock(self, symbol, filters):
        """Scan a single stock for trading opportunities"""
        try:
            data = self.get_stock_data(symbol, period="3mo")
            if data is None:
                return None

            # Calculate indicators
            close = data['Close']
            rsi = self.calculate_rsi(close, period=14)
            sma_20 = close.rolling(window=20).mean()
            sma_50 = close.rolling(window=50).mean()

            # Current values
            current_close = close.iloc[-1]
            current_rsi = rsi.iloc[-1]
            current_volume = data['Volume'].iloc[-1]
            avg_volume = data['Volume'].tail(20).mean()
            volume_ratio = current_volume / avg_volume if avg_volume > 0 else 0

            # Filters
            rsi_min = filters.get('rsi_min', 30)
            rsi_max = filters.get('rsi_max', 75)
            vol_ratio_min = filters.get('min_volume_ratio', 1.2)

            # Check criteria
            if not (rsi_min <= current_rsi <= rsi_max):
                return None

            if volume_ratio < vol_ratio_min:
                return None

            # Check support
            if sma_20.iloc[-1] is None or current_close < sma_20.iloc[-1] * 0.97:
                return None

            # Calculate pattern strength (simple score)
            strength = 50

            # RSI positioning (sweet spot is 40-60)
            if 40 <= current_rsi <= 60:
                strength += 15
            elif 35 <= current_rsi <= 65:
                strength += 10

            # Volume contribution
            if volume_ratio >= 2.0:
                strength += 20
            elif volume_ratio >= 1.5:
                strength += 10

            # Price action
            if current_close > sma_20.iloc[-1]:
                strength += 15

            # Momentum
            price_change = (current_close - close.iloc[-5]) / close.iloc[-5] * 100
            if price_change > 0:
                strength += min(10, price_change / 2)

            strength = min(100, max(0, strength))

            if strength < filters.get('min_strength', 65):
                return None

            return {
                'symbol': symbol.replace('.NS', ''),
                'price': current_close,
                'rsi': current_rsi,
                'sma_20': sma_20.iloc[-1],
                'sma_50': sma_50.iloc[-1],
                'volume_ratio': volume_ratio,
                'strength': strength,
                'date': data.index[-1].strftime('%Y-%m-%d'),
                'change_5d': price_change
            }

        except Exception as e:
            logger.debug(f"Error scanning {symbol}: {e}")
            return None

    def run_scan(self, num_stocks=100, min_strength=65):
        """Run scanning on multiple stocks"""
        filters = {
            'rsi_min': 30,
            'rsi_max': 75,
            'min_volume_ratio': 1.2,
            'min_strength': min_strength
        }

        stocks_to_scan = COMPLETE_NSE_FO_UNIVERSE[:num_stocks]
        results = []

        logger.info(f"🚀 Starting scan of {len(stocks_to_scan)} stocks")
        logger.info(f"   Pattern Strength: {min_strength}%+")
        logger.info(f"   Time: {datetime.now(self.ist).strftime('%H:%M %Z')}\n")

        # Send initial Telegram message
        if self.telegram_enabled:
            self.send_telegram(
                f"🚀 *NSE Stock Scanner Started*\n"
                f"Scanning {len(stocks_to_scan)} stocks\n"
                f"Time: {datetime.now(self.ist).strftime('%H:%M %Z')}\n"
                f"Pattern Strength: {min_strength}%+"
            )

        # Scan stocks with threading
        with ThreadPoolExecutor(max_workers=5) as executor:
            futures = {
                executor.submit(self.scan_stock, symbol, filters): symbol
                for symbol in stocks_to_scan
            }

            completed = 0
            for future in as_completed(futures):
                try:
                    result = future.result()
                    if result:
                        results.append(result)
                        logger.info(f"✅ Found: {result['symbol']} - Strength: {result['strength']:.0f}%")
                    completed += 1

                    if completed % 20 == 0:
                        logger.info(f"📊 Progress: {completed}/{len(stocks_to_scan)}")

                except Exception as e:
                    logger.error(f"Scan error: {e}")
                    completed += 1

        logger.info(f"\n{'='*50}")
        logger.info(f"✅ SCAN COMPLETE - Found {len(results)} stocks")
        logger.info(f"{'='*50}\n")

        return results

    def export_to_excel(self, results):
        """Export results to Excel"""
        if not results:
            logger.warning("No results to export")
            return None

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        timestamp = datetime.now(self.ist).strftime('%Y%m%d_%H%M%S')
        filename = f"stock_scan_results_{timestamp}.xlsx"

        # Create DataFrame
        df = pd.DataFrame(results)

        # Export to Excel
        df.to_excel(filename, index=False)

        # Format Excel file
        wb = openpyxl.load_workbook(filename)
        ws = wb.active

        # Header formatting
        header_fill = PatternFill(start_color="ff6b00", end_color="ff6b00", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Auto-adjust columns
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filename)
        logger.info(f"📊 Results exported to: {filename}")

        return filename

    def format_telegram_message(self, results):
        """Format results for Telegram"""
        if not results:
            return "❌ No stocks found meeting the criteria."

        # Sort by strength
        results_sorted = sorted(results, key=lambda x: x['strength'], reverse=True)

        message = f"✅ *NSE Stock Scanner Results*\n"
        message += f"Time: {datetime.now(self.ist).strftime('%H:%M %Z')}\n"
        message += f"Stocks Found: {len(results)}\n\n"

        message += "*Top Candidates:*\n"
        for r in results_sorted[:10]:
            message += (
                f"• *{r['symbol']}*: ₹{r['price']:.2f} | "
                f"Str:{r['strength']:.0f}% | RSI:{r['rsi']:.0f} | "
                f"Vol:{r['volume_ratio']:.1f}x\n"
            )

        if len(results) > 10:
            message += f"\n... and {len(results) - 10} more stocks\n"

        message += f"\n📊 See Excel file for complete analysis"

        return message


def main():
    import argparse

    parser = argparse.ArgumentParser(description='NSE Stock Scanner')
    parser.add_argument('--stocks', type=int, default=100, help='Number of stocks to scan')
    parser.add_argument('--strength', type=int, default=65, help='Minimum pattern strength')
    parser.add_argument('--no-telegram', action='store_true', help='Skip Telegram')

    args = parser.parse_args()

    scanner = SimpleStockScanner()

    if args.no_telegram:
        scanner.telegram_enabled = False

    # Run scan
    results = scanner.run_scan(num_stocks=args.stocks, min_strength=args.strength)

    # Export to Excel
    excel_file = scanner.export_to_excel(results)

    # Send to Telegram
    if scanner.telegram_enabled and results:
        message = scanner.format_telegram_message(results)
        scanner.send_telegram(message)

    print(f"\n{'='*60}")
    print(f"SCAN RESULTS SUMMARY")
    print(f"{'='*60}")
    print(f"Total Stocks Found: {len(results)}")
    print(f"Results Saved: {excel_file}")
    print(f"{'='*60}\n")

    if results:
        print("Top 5 Stocks by Strength:")
        for i, r in enumerate(sorted(results, key=lambda x: x['strength'], reverse=True)[:5], 1):
            print(f"{i}. {r['symbol']:12} - Strength: {r['strength']:5.0f}% | Price: ₹{r['price']:8.2f} | RSI: {r['rsi']:6.1f}")

    return results


if __name__ == "__main__":
    main()
